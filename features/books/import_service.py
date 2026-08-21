"""추천도서 일괄 등록. preview는 DB를 쓰지 않고, confirm만 트랜잭션으로 반영한다."""
from __future__ import annotations

import csv
import io
import json
import tempfile
import unicodedata
import uuid
from pathlib import Path

from extensions import db
from feature_models import (
    Book,
    clean_book_author,
    clean_book_title,
    normalize_book_title,
)
from features.books.service import create_book_record
from features.reading.classify import VALID_GRADE_BANDS, normalize_grade_band

ACTION_CREATE = 'create'
ACTION_REUSE = 'reuse'
ACTION_REVIEW = 'review'
ACTION_ERROR = 'error'
MAX_IMPORT_ROWS = 400
PREVIEW_PREFIX = 'clc_rec_import_'

TITLE_HEADERS = {'title', '제목', '책제목', '도서명', '책이름'}
AUTHOR_HEADERS = {'author', '저자', '글쓴이', '작가'}
GRADE_HEADERS = {'grade_band', 'gradeband', '학년군', '학년밴드', '대상학년'}


class RecommendedImportError(Exception):
    def __init__(self, message, code='error'):
        super().__init__(message)
        self.message = message
        self.code = code


def normalize_author_key(author):
    if author is None:
        return ''
    text = unicodedata.normalize('NFKC', str(author)).strip()
    text = ' '.join(text.split())
    return text.casefold()


def _authors_clearly_match(left, right):
    left_key = normalize_author_key(left)
    right_key = normalize_author_key(right)
    if not left_key and not right_key:
        return True
    if not left_key or not right_key:
        return False
    return left_key == right_key


def _authors_clearly_different(left, right):
    left_key = normalize_author_key(left)
    right_key = normalize_author_key(right)
    return bool(left_key and right_key and left_key != right_key)


def _header_key(raw):
    text = unicodedata.normalize('NFKC', '' if raw is None else str(raw)).strip().casefold()
    return text.replace(' ', '').replace('_', '')


def _pick_column(headers, aliases):
    for index, header in enumerate(headers):
        key = _header_key(header)
        if key in aliases or header.strip().casefold() in aliases:
            return index
    return None


def parse_paste_text(raw_text):
    rows = []
    text = '' if raw_text is None else str(raw_text)
    for line_no, raw_line in enumerate(text.splitlines(), start=1):
        line = raw_line.strip()
        if not line:
            continue
        if '|' in line:
            parts = [part.strip() for part in line.split('|', 1)]
        elif '\t' in line:
            parts = [part.strip() for part in line.split('\t', 1)]
        else:
            parts = [line]
        title = parts[0] if parts else ''
        author = parts[1] if len(parts) > 1 else None
        rows.append({
            'line_no': line_no,
            'title': title,
            'author': author,
            'grade_band': None,
        })
    return rows


def parse_csv_bytes(raw_bytes):
    if not raw_bytes:
        return []
    text = raw_bytes.decode('utf-8-sig')
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=',\t;|')
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    return _rows_from_tabular(reader)


def parse_xlsx_bytes(raw_bytes):
    # TODO(security-hardening): 파일 크기, ZIP bomb, MIME/매직바이트 검증은
    # 전체 핵심 기능 완료 후 최종 보안 단계에서 별도로 진행한다.
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(raw_bytes), read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = []
        for row in sheet.iter_rows(values_only=True):
            rows.append([('' if cell is None else str(cell)) for cell in row])
    finally:
        workbook.close()
    return _rows_from_tabular(rows)


def _rows_from_tabular(table):
    rows_iter = iter(table)
    try:
        header = next(rows_iter)
    except StopIteration:
        return []
    headers = [('' if cell is None else str(cell)) for cell in header]
    title_idx = _pick_column(headers, TITLE_HEADERS)
    author_idx = _pick_column(headers, AUTHOR_HEADERS)
    grade_idx = _pick_column(headers, GRADE_HEADERS)
    if title_idx is None and len(headers) >= 1:
        # 헤더 없이 제목,저자 만 있는 파일도 허용
        first_keys = {_header_key(cell) for cell in headers}
        looks_like_header = bool(first_keys & (TITLE_HEADERS | AUTHOR_HEADERS | GRADE_HEADERS))
        if not looks_like_header:
            data_rows = [headers] + [list(row) for row in rows_iter]
            parsed = []
            for line_no, row in enumerate(data_rows, start=1):
                values = [('' if cell is None else str(cell)).strip() for cell in row]
                parsed.append({
                    'line_no': line_no,
                    'title': values[0] if values else '',
                    'author': values[1] if len(values) > 1 else None,
                    'grade_band': values[2] if len(values) > 2 else None,
                })
            return parsed
        title_idx = 0

    parsed = []
    for offset, row in enumerate(rows_iter, start=2):
        values = [('' if cell is None else str(cell)).strip() for cell in row]
        if not any(values):
            continue
        title = values[title_idx] if title_idx is not None and title_idx < len(values) else ''
        author = values[author_idx] if author_idx is not None and author_idx < len(values) else None
        grade_band = values[grade_idx] if grade_idx is not None and grade_idx < len(values) else None
        parsed.append({
            'line_no': offset,
            'title': title,
            'author': author,
            'grade_band': grade_band,
        })
    return parsed


def _match_existing_book(title, author, extra_candidates=None):
    cleaned_title = clean_book_title(title)
    key = normalize_book_title(cleaned_title)
    if not key:
        return ACTION_ERROR, None, '제목은 필수입니다.'
    candidates = (
        Book.query.filter(Book.normalized_key == key)
        .order_by(Book.id.asc())
        .all()
    )
    extras = []
    for item in extra_candidates or []:
        if normalize_book_title(getattr(item, 'title', None) or '') == key:
            extras.append(item)
    combined = list(candidates) + extras
    if not combined:
        return ACTION_CREATE, None, None
    if len(combined) != 1:
        return ACTION_REVIEW, None, '같은 제목의 책이 여러 권 있어 확인이 필요합니다.'

    candidate = combined[0]
    if _authors_clearly_match(candidate.author, author):
        if getattr(candidate, 'is_challenge_eligible', False):
            return ACTION_ERROR, candidate, '이미 도전도서인 책은 추천도서로 지정할 수 없습니다.'
        return ACTION_REUSE, candidate, None
    if _authors_clearly_different(candidate.author, author):
        return ACTION_CREATE, None, None
    return ACTION_REVIEW, candidate, '제목은 같지만 저자가 불명확하여 자동 연결하지 않습니다.'


class _PendingBook:
    def __init__(self, title, author):
        self.id = None
        self.title = title
        self.author = author


def preview_rows(raw_rows, default_grade_band=None):
    default_band = normalize_grade_band(default_grade_band)
    preview = []
    counts = {
        'total': 0,
        ACTION_CREATE: 0,
        ACTION_REUSE: 0,
        ACTION_REVIEW: 0,
        ACTION_ERROR: 0,
    }
    pending = []
    for raw in raw_rows:
        counts['total'] += 1
        title = clean_book_title(raw.get('title'))
        author = clean_book_author(raw.get('author'))
        row_band = normalize_grade_band(raw.get('grade_band'))
        if raw.get('grade_band') not in (None, '') and row_band is None:
            action = ACTION_ERROR
            message = '학년군이 올바르지 않습니다. 2-3 또는 4-6만 사용할 수 있습니다.'
            grade_band = None
            book = None
        elif not title:
            action = ACTION_ERROR
            message = '제목은 필수입니다.'
            grade_band = row_band or default_band
            book = None
        else:
            grade_band = row_band or default_band
            if grade_band not in VALID_GRADE_BANDS:
                action = ACTION_ERROR
                message = '학년군을 선택하거나 파일에 grade_band를 넣어 주세요.'
                book = None
            else:
                action, book, message = _match_existing_book(title, author, extra_candidates=pending)
                if action == ACTION_CREATE:
                    pending.append(_PendingBook(title, author))

        counts[action] = counts.get(action, 0) + 1
        line_no = raw.get('line_no')
        row_key = str(line_no if line_no is not None else f'i{counts["total"]}')
        preview.append({
            'row_key': row_key,
            'line_no': line_no,
            'title': title or ('' if raw.get('title') is None else str(raw.get('title')).strip()),
            'author': author,
            'grade_band': grade_band,
            'action': action,
            'message': message,
            'book_id': None if book is None else book.id,
            'book_title': None if book is None else book.title,
            'book_author': None if book is None else book.author,
            'selectable': action in {ACTION_CREATE, ACTION_REUSE},
        })
    return {
        'counts': {
            'total': counts['total'],
            'create': counts[ACTION_CREATE],
            'reuse': counts[ACTION_REUSE],
            'review': counts[ACTION_REVIEW],
            'error': counts[ACTION_ERROR],
        },
        'rows': preview,
        'default_grade_band': default_band,
    }


def store_preview(preview, extra=None):
    token = uuid.uuid4().hex
    payload = {
        'token': token,
        'preview': preview,
        'extra': extra or {},
    }
    path = Path(tempfile.gettempdir()) / f'{PREVIEW_PREFIX}{token}.json'
    path.write_text(json.dumps(payload, ensure_ascii=False), encoding='utf-8')
    return token


def load_preview(token):
    if not token:
        raise RecommendedImportError('미리보기 정보가 없습니다. 다시 올려 주세요.', code='preview_missing')
    safe = ''.join(ch for ch in str(token) if ch.isalnum())
    if not safe or safe != str(token):
        raise RecommendedImportError('미리보기 정보가 올바르지 않습니다.', code='preview_invalid')
    path = Path(tempfile.gettempdir()) / f'{PREVIEW_PREFIX}{safe}.json'
    if not path.exists():
        raise RecommendedImportError('미리보기가 만료되었습니다. 다시 올려 주세요.', code='preview_expired')
    payload = json.loads(path.read_text(encoding='utf-8'))
    return payload


def clear_preview(token):
    try:
        safe = ''.join(ch for ch in str(token) if ch.isalnum())
        path = Path(tempfile.gettempdir()) / f'{PREVIEW_PREFIX}{safe}.json'
        if path.exists():
            path.unlink()
    except OSError:
        pass


def apply_preview(preview, selected_row_keys=None):
    """create/reuse 만 반영. review/error 와 선택 해제 행은 건너뛴다. 실패 시 전부 롤백.

    selected_row_keys가 None이면 적용 가능 행을 모두 반영한다(서비스/기존 테스트).
    HTTP confirm은 체크된 row_key만 넘긴다.
    """
    rows = (preview or {}).get('rows') or []
    created = 0
    reused = 0
    skipped = 0
    selected = None if selected_row_keys is None else {str(key) for key in selected_row_keys}
    try:
        for row in rows:
            action = row.get('action')
            row_key = str(row.get('row_key') or row.get('line_no') or '')
            if selected is not None and row_key not in selected:
                skipped += 1
                continue
            if action not in {ACTION_CREATE, ACTION_REUSE}:
                skipped += 1
                continue
            title = row.get('title')
            author = row.get('author')
            grade_band = normalize_grade_band(row.get('grade_band'))
            if grade_band not in VALID_GRADE_BANDS:
                raise RecommendedImportError('학년군이 올바르지 않습니다.', code='invalid_grade_band')

            current_action, existing, _message = _match_existing_book(title, author)
            if current_action == ACTION_REVIEW:
                skipped += 1
                continue
            if current_action == ACTION_ERROR:
                raise RecommendedImportError('제목은 필수입니다.', code='title_required')

            if current_action == ACTION_REUSE and existing is not None:
                if existing.is_challenge_eligible:
                    raise RecommendedImportError(
                        '이미 도전도서인 책은 추천도서로 지정할 수 없습니다.',
                        code='challenge_conflict',
                    )
                existing.is_recommended = True
                existing.grade_band = grade_band
                reused += 1
                continue

            book, _similar = create_book_record(title, author, commit=False)
            book.is_recommended = True
            book.grade_band = grade_band
            created += 1

        db.session.commit()
    except Exception:
        db.session.rollback()
        raise

    return {
        'created': created,
        'reused': reused,
        'skipped': skipped,
    }
