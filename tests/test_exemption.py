"""Step 6: 5~6학년 추천독서 보상 선택과 학습 면제권 장부."""
from __future__ import annotations

import json
import tempfile
import unittest
from datetime import date, timedelta
from pathlib import Path
from unittest import mock

from sqlalchemy.exc import IntegrityError

from tests.helpers import bootstrap_test_app

app, db = bootstrap_test_app()

from app import Child, DailyPoints, PointsHistory, User, get_backup_data  # noqa: E402
from feature_models import (  # noqa: E402
    Book,
    ChildReading,
    EVENT_RECOMMENDED_COMPLETE,
    EVENT_RECOMMENDED_START,
    ExemptionTicket,
    ExemptionTicketSource,
    ExemptionUsage,
    LearningProgressEntry,
    LearningSubject,
    ReadingDay,
    ReadingRewardEvent,
    STATUS_COMPLETED,
)
from features.books.restore import restore_books_from_backup_data  # noqa: E402
from features.books.service import create_book_record, update_recommended_flags  # noqa: E402
from features.dates import kst_today  # noqa: E402
from features.exemption.policy import next_issue_on, ticket_expires_on  # noqa: E402
from features.exemption.restore import restore_exemptions_from_backup_data  # noqa: E402
from features.exemption.service import (  # noqa: E402
    ExemptionError,
    child_exemption_snapshot,
    issue_exemption_ticket,
    revoke_exemption_ticket,
    set_reward_mode,
    use_exemption_ticket,
)
from features.progress.restore import restore_learning_progress_from_backup_data  # noqa: E402
from features.progress.service import (  # noqa: E402
    ensure_default_subjects,
    save_progress_entry,
    set_subject_active,
)
from features.reading.restore import restore_readings_from_backup_data  # noqa: E402
from features.reading.rewards import RewardError, approve_recommended_reward  # noqa: E402
from features.reading.service import ReadingError, complete_current, start_book  # noqa: E402
from features.reading.session import SESSION_CHILD_ID, SESSION_CHILD_SLUG, SESSION_VERIFIED_AT  # noqa: E402
from features.reading.policy import now_utc  # noqa: E402


class ExemptionLedgerTests(unittest.TestCase):
    def setUp(self):
        self.ctx = app.app_context()
        self.ctx.push()
        db.session.remove()
        db.drop_all()
        db.create_all()
        self.teacher = User(
            username='ex_teacher',
            name='면제교사',
            role='돌봄선생님',
            email='ex-teacher@example.test',
            password_hash='',
        )
        self.developer = User(
            username='ex_dev',
            name='면제개발자',
            role='개발자',
            email='ex-dev@example.test',
            password_hash='',
        )
        self.viewer = User(
            username='ex_viewer',
            name='면제열람',
            role='학생열람',
            email='studentview-ex@example.test',
            password_hash='',
        )
        db.session.add_all([self.teacher, self.developer, self.viewer])
        db.session.commit()
        self.teacher_id = self.teacher.id
        self.developer_id = self.developer.id
        self.viewer_id = self.viewer.id
        self.today = date(2026, 8, 20)
        self.client = app.test_client()
        self._slug_n = 0
        self.book_n = 0
        ensure_default_subjects()
        self.math = LearningSubject.query.filter_by(key='math').one()
        self.korean = LearningSubject.query.filter_by(key='korean').one()
        self.ssen = LearningSubject.query.filter_by(key='ssen').one()

    def tearDown(self):
        db.session.remove()
        self.ctx.pop()

    def _child(self, name, grade):
        self._slug_n += 1
        child = Child(name=name, grade=grade, viewer_slug=f'{self._slug_n:024x}')
        db.session.add(child)
        db.session.commit()
        return child

    def _rec_book(self, title=None):
        self.book_n += 1
        book, _ = create_book_record(title or f'추천{self.book_n}', '면제작가')
        update_recommended_flags(book, is_recommended=True, grade_band='4-6')
        return book

    def _login(self, user, client=None):
        client = client or self.client
        user_id = user if isinstance(user, int) else user.id
        with client.session_transaction() as sess:
            sess['_user_id'] = str(user_id)
            sess['_fresh'] = True
        return client

    def _verify(self, child, client=None):
        client = client or self.client
        with client.session_transaction() as sess:
            sess[SESSION_CHILD_ID] = int(child.id)
            sess[SESSION_CHILD_SLUG] = child.viewer_slug
            sess[SESSION_VERIFIED_AT] = now_utc().isoformat()

    def _start(self, child, book, activity_date):
        return start_book(
            child.id,
            self.teacher.id,
            'teacher',
            activity_date=activity_date,
            book_id=book.id,
        )

    def _complete_exemption(self, child, book, activity_date):
        reading, _ = self._start(child, book, activity_date)
        set_reward_mode(reading, 'exemption', self.teacher)
        complete_current(child.id, self.teacher.id, 'teacher', activity_date=activity_date)
        return ChildReading.query.get(reading.id)

    def _complete_points(self, child, book, activity_date):
        reading, _ = self._start(child, book, activity_date)
        set_reward_mode(reading, 'points', self.teacher)
        complete_current(child.id, self.teacher.id, 'teacher', activity_date=activity_date)
        return ChildReading.query.get(reading.id)

    def test_date_boundaries(self):
        issued = date(2026, 8, 20)
        self.assertEqual(ticket_expires_on(issued), date(2026, 9, 2))
        self.assertEqual(next_issue_on(issued), date(2026, 9, 3))

    def test_grade_5_6_can_choose_points_and_exemption(self):
        for grade in (5, 6):
            child = self._child(f'선택{grade}', grade)
            reading, _ = self._start(child, self._rec_book(), self.today)
            self.assertIsNone(reading.reward_mode)
            set_reward_mode(reading, 'points', self.teacher)
            self.assertEqual(ChildReading.query.get(reading.id).reward_mode, 'points')
            set_reward_mode(reading, 'exemption', self.teacher)
            self.assertEqual(ChildReading.query.get(reading.id).reward_mode, 'exemption')

    def test_start_without_mode_is_allowed(self):
        child = self._child('미선택시작', 5)
        reading, _ = self._start(child, self._rec_book(), self.today)
        self.assertIsNone(reading.reward_mode)
        self.assertEqual(reading.status, 'in_progress')
        self.assertEqual(ExemptionTicket.query.count(), 0)
        self.assertEqual(ReadingRewardEvent.query.count(), 0)

    def test_viewer_can_set_mode(self):
        child = self._child('열람선택', 6)
        reading, _ = self._start(child, self._rec_book(), self.today)
        self._login(self.viewer_id)
        self._verify(child)
        resp = self.client.post(
            f'/viewer/report/{child.viewer_slug}/reading/reward-mode',
            data={'child_reading_id': str(reading.id), 'reward_mode': 'exemption'},
            follow_redirects=False,
        )
        self.assertEqual(resp.status_code, 302)
        db.session.expire_all()
        self.assertEqual(ChildReading.query.get(reading.id).reward_mode, 'exemption')

    def test_teacher_can_set_mode(self):
        child = self._child('교사선택', 5)
        reading, _ = self._start(child, self._rec_book(), self.today)
        self._login(self.teacher_id)
        resp = self.client.post(
            f'/children/{child.id}/reading/reward-mode',
            data={'child_reading_id': str(reading.id), 'reward_mode': 'points'},
            follow_redirects=False,
        )
        self.assertEqual(resp.status_code, 302)
        db.session.expire_all()
        self.assertEqual(ChildReading.query.get(reading.id).reward_mode, 'points')

    def test_viewer_cannot_approve_or_issue_or_use(self):
        child = self._child('열람차단', 5)
        reading = self._complete_exemption(child, self._rec_book(), self.today)
        self._login(self.viewer_id)
        self._verify(child)
        approve = self.client.post(
            f'/children/{child.id}/reading/reward/approve',
            data={'child_reading_id': str(reading.id), 'event_type': EVENT_RECOMMENDED_START},
            follow_redirects=False,
        )
        issue = self.client.post(
            f'/children/{child.id}/exemption/issue',
            follow_redirects=False,
        )
        self.assertEqual(approve.status_code, 302)
        self.assertIn('/viewer', approve.headers.get('Location', ''))
        self.assertEqual(issue.status_code, 302)
        self.assertIn('/viewer', issue.headers.get('Location', ''))
        self.assertEqual(ReadingRewardEvent.query.count(), 0)
        self.assertEqual(ExemptionTicket.query.count(), 0)
        self._login(self.teacher_id)
        with mock.patch('features.exemption.service.kst_today', return_value=self.today):
            issue_exemption_ticket(child.id, self.teacher, today=self.today)
            ticket = ExemptionTicket.query.one()
            self._login(self.viewer_id)
            self._verify(child)
            use = self.client.post(
                f'/children/{child.id}/exemption/{ticket.id}/use',
                data={'subject_key': 'korean', 'used_on': self.today.isoformat()},
                follow_redirects=False,
            )
        self.assertEqual(use.status_code, 302)
        self.assertIn('/viewer', use.headers.get('Location', ''))
        self.assertEqual(ExemptionTicket.query.one().status, 'active')
        self.assertEqual(ExemptionUsage.query.count(), 0)

    def test_points_mode_start_100_complete_200(self):
        for grade in (5, 6):
            child = self._child(f'포인트{grade}', grade)
            reading, _ = self._start(child, self._rec_book(), self.today)
            set_reward_mode(reading, 'points', self.teacher)
            approve_recommended_reward(reading, self.teacher, EVENT_RECOMMENDED_START)
            complete_current(child.id, self.teacher.id, 'teacher', activity_date=self.today)
            reading = ChildReading.query.get(reading.id)
            approve_recommended_reward(reading, self.teacher, EVENT_RECOMMENDED_COMPLETE)
            daily = DailyPoints.query.filter_by(child_id=child.id).one()
            self.assertEqual(daily.reading_points, 0)
            self.assertEqual(daily.manual_points, 300)
            self.assertEqual(Child.query.get(child.id).cumulative_points, 300)

    def test_exemption_and_null_block_point_rewards(self):
        child = self._child('면제포인트차단', 5)
        reading, _ = self._start(child, self._rec_book(), self.today)
        with self.assertRaises(RewardError) as unset_err:
            approve_recommended_reward(reading, self.teacher, EVENT_RECOMMENDED_START)
        self.assertEqual(unset_err.exception.code, 'mode_unset')
        set_reward_mode(reading, 'exemption', self.teacher)
        with self.assertRaises(RewardError) as ex_err:
            approve_recommended_reward(reading, self.teacher, EVENT_RECOMMENDED_START)
        self.assertEqual(ex_err.exception.code, 'exemption_mode')
        complete_current(child.id, self.teacher.id, 'teacher', activity_date=self.today)
        reading = ChildReading.query.get(reading.id)
        with self.assertRaises(RewardError):
            approve_recommended_reward(reading, self.teacher, EVENT_RECOMMENDED_COMPLETE)
        self.assertEqual(DailyPoints.query.count(), 0)
        self.assertEqual(ReadingRewardEvent.query.count(), 0)

    def test_mode_change_free_before_benefits(self):
        child = self._child('자유변경', 5)
        reading, _ = self._start(child, self._rec_book(), self.today)
        set_reward_mode(reading, 'points', self.teacher)
        set_reward_mode(reading, 'exemption', self.teacher)
        set_reward_mode(reading, 'points', self.teacher)
        self.assertEqual(ChildReading.query.get(reading.id).reward_mode, 'points')

    def test_cannot_switch_to_exemption_after_points_awarded(self):
        child = self._child('이중방지P', 5)
        reading, _ = self._start(child, self._rec_book(), self.today)
        set_reward_mode(reading, 'points', self.teacher)
        approve_recommended_reward(reading, self.teacher, EVENT_RECOMMENDED_START)
        with self.assertRaises(ExemptionError) as err:
            set_reward_mode(reading, 'exemption', self.teacher)
        self.assertEqual(err.exception.code, 'points_already_awarded')
        from features.reading.rewards import revoke_recommended_reward
        revoke_recommended_reward(reading, self.teacher, EVENT_RECOMMENDED_START)
        set_reward_mode(reading, 'exemption', self.teacher)
        self.assertEqual(ChildReading.query.get(reading.id).reward_mode, 'exemption')

    def test_first_completion_makes_ticket_eligible_but_not_auto_issued(self):
        child = self._child('첫권', 5)
        reading = self._complete_exemption(child, self._rec_book(), self.today)
        snap = child_exemption_snapshot(child.id, self.today)
        self.assertTrue(snap['condition_met'])
        self.assertTrue(snap['can_issue'])
        self.assertEqual(snap['needed'], 1)
        self.assertEqual(ExemptionTicket.query.count(), 0)
        self.assertEqual(reading.status, STATUS_COMPLETED)

    def test_teacher_issues_first_ticket_with_one_source(self):
        child = self._child('발급1', 5)
        reading = self._complete_exemption(child, self._rec_book(), self.today)
        ticket = issue_exemption_ticket(child.id, self.teacher, today=self.today)
        self.assertEqual(ticket.status, 'active')
        self.assertEqual(ticket.issued_on, self.today)
        self.assertEqual(ticket.expires_on, date(2026, 9, 2))
        sources = ExemptionTicketSource.query.filter_by(exemption_ticket_id=ticket.id).all()
        self.assertEqual([row.child_reading_id for row in sources], [reading.id])

    def test_second_book_alone_not_enough_third_book_qualifies(self):
        child = self._child('패턴', 5)
        a = self._complete_exemption(child, self._rec_book('권A'), self.today)
        ticket1 = issue_exemption_ticket(child.id, self.teacher, today=self.today)
        b = self._complete_exemption(child, self._rec_book('권B'), self.today + timedelta(days=1))
        snap = child_exemption_snapshot(child.id, self.today + timedelta(days=1))
        self.assertFalse(snap['condition_met'])
        self.assertEqual(snap['have'], 1)
        self.assertEqual(snap['needed'], 2)
        with self.assertRaises(ExemptionError) as err:
            issue_exemption_ticket(child.id, self.teacher, today=self.today + timedelta(days=14))
        self.assertEqual(err.exception.code, 'not_ready')
        c = self._complete_exemption(child, self._rec_book('권C'), self.today + timedelta(days=2))
        day = date(2026, 9, 3)
        ticket2 = issue_exemption_ticket(child.id, self.teacher, today=day)
        ids = {
            row.child_reading_id
            for row in ExemptionTicketSource.query.filter_by(exemption_ticket_id=ticket2.id)
        }
        self.assertEqual(ids, {b.id, c.id})
        self.assertNotIn(a.id, ids)
        self.assertEqual(ticket1.id, ExemptionTicket.query.order_by(ExemptionTicket.id.asc()).first().id)

    def test_oldest_unconsumed_selected_and_no_double_consume(self):
        child = self._child('오래된순', 6)
        a = self._complete_exemption(child, self._rec_book('오래A'), self.today)
        issue_exemption_ticket(child.id, self.teacher, today=self.today)
        b = self._complete_exemption(child, self._rec_book('오래B'), self.today + timedelta(days=1))
        c = self._complete_exemption(child, self._rec_book('오래C'), self.today + timedelta(days=2))
        d = self._complete_exemption(child, self._rec_book('오래D'), self.today + timedelta(days=3))
        ticket2 = issue_exemption_ticket(child.id, self.teacher, today=date(2026, 9, 3))
        ids = [
            row.child_reading_id
            for row in ExemptionTicketSource.query.filter_by(exemption_ticket_id=ticket2.id)
            .order_by(ExemptionTicketSource.id.asc())
        ]
        self.assertEqual(ids, [b.id, c.id])
        self.assertNotIn(a.id, ids)
        self.assertNotIn(d.id, ids)

    def test_max_active_one_and_double_issue(self):
        child = self._child('한방', 5)
        self._complete_exemption(child, self._rec_book(), self.today)
        first = issue_exemption_ticket(child.id, self.teacher, today=self.today)
        with self.assertRaises(ExemptionError) as err:
            issue_exemption_ticket(child.id, self.teacher, today=self.today)
        self.assertEqual(err.exception.code, 'already_holding')
        db.session.add(ExemptionTicket(
            child_id=child.id,
            issued_on=self.today,
            expires_on=ticket_expires_on(self.today),
            status='active',
            policy_version='v1',
            issued_by_user_id=self.teacher.id,
        ))
        with self.assertRaises(IntegrityError):
            db.session.commit()
        db.session.rollback()
        self.assertEqual(ExemptionTicket.query.filter_by(status='active').count(), 1)
        self.assertEqual(first.status, 'active')

    def test_progress_preserved_during_hold_and_cooldown(self):
        child = self._child('보존', 5)
        self._complete_exemption(child, self._rec_book('보A'), self.today)
        ticket = issue_exemption_ticket(child.id, self.teacher, today=self.today)
        self._complete_exemption(child, self._rec_book('보B'), self.today + timedelta(days=1))
        self._complete_exemption(child, self._rec_book('보C'), self.today + timedelta(days=2))
        holding = child_exemption_snapshot(child.id, self.today + timedelta(days=2))
        self.assertTrue(holding['condition_met'])
        self.assertTrue(holding['holding_blocks'])
        self.assertFalse(holding['can_issue'])
        use_exemption_ticket(ticket.id, 'math', self.teacher, used_on=self.today + timedelta(days=1), today=self.today + timedelta(days=1))
        cooldown = child_exemption_snapshot(child.id, self.today + timedelta(days=2))
        self.assertTrue(cooldown['condition_met'])
        self.assertTrue(cooldown['cooldown_blocks'])
        self.assertEqual(cooldown['next_issue_on'], date(2026, 9, 3))
        self.assertFalse(cooldown['can_issue'])
        self.assertEqual(cooldown['unconsumed_count'], 2)

    def test_use_valid_range_and_reject_day_15(self):
        child = self._child('유효', 5)
        self._complete_exemption(child, self._rec_book(), self.today)
        ticket = issue_exemption_ticket(child.id, self.teacher, today=self.today)
        use_exemption_ticket(ticket.id, 'math', self.teacher, used_on=date(2026, 9, 2), today=date(2026, 9, 2))
        self.assertEqual(ExemptionTicket.query.get(ticket.id).status, 'used')
        child2 = self._child('만료사용', 5)
        self._complete_exemption(child2, self._rec_book(), self.today)
        ticket2 = issue_exemption_ticket(child2.id, self.teacher, today=self.today)
        with self.assertRaises(ExemptionError) as err:
            use_exemption_ticket(ticket2.id, 'math', self.teacher, used_on=date(2026, 9, 3), today=date(2026, 9, 3))
        self.assertEqual(err.exception.code, 'expired')

    def test_backfill_used_on_inside_window(self):
        child = self._child('후기록', 5)
        self._complete_exemption(child, self._rec_book(), self.today)
        ticket = issue_exemption_ticket(child.id, self.teacher, today=self.today)
        use_exemption_ticket(ticket.id, 'korean', self.teacher, used_on=date(2026, 8, 21), today=date(2026, 8, 25))
        usage = ExemptionUsage.query.one()
        self.assertEqual(usage.used_on, date(2026, 8, 21))
        with self.assertRaises(ExemptionError):
            use_exemption_ticket(ticket.id, 'math', self.teacher, used_on=date(2026, 8, 22), today=date(2026, 8, 25))

    def test_subject_validation_and_past_usage_kept(self):
        child = self._child('과목검증', 5)
        self._complete_exemption(child, self._rec_book(), self.today)
        ticket = issue_exemption_ticket(child.id, self.teacher, today=self.today)
        for raw in (None, '', 'english'):
            with self.assertRaises(ExemptionError) as err:
                use_exemption_ticket(ticket.id, raw, self.teacher, today=self.today)
            self.assertEqual(err.exception.code, 'subject_not_allowed')
            db.session.refresh(ticket)
            self.assertEqual(ticket.status, 'active')
            self.assertEqual(ExemptionUsage.query.count(), 0)
        use_exemption_ticket(ticket.id, 'math', self.teacher, today=self.today)
        usage = ExemptionUsage.query.one()
        self.assertEqual(usage.subject_key, 'math')
        self.assertEqual(usage.subject_name, '수학')
        LearningSubject.query.delete()
        db.session.commit()
        self.assertEqual(ExemptionUsage.query.one().subject_key, 'math')
        self.assertEqual(ExemptionUsage.query.one().subject_name, '수학')
        self.assertEqual(ExemptionUsage.query.count(), 1)

    def test_http_use_requires_explicit_subject_no_korean_fallback(self):
        self._login(self.teacher_id)

        def issue_one(name):
            child = self._child(name, 5)
            self._complete_exemption(child, self._rec_book(), self.today)
            ticket = issue_exemption_ticket(child.id, self.teacher, today=self.today)
            return child, ticket

        child, ticket = issue_one('누락과목')
        with mock.patch('features.exemption.service.kst_today', return_value=self.today):
            missing = self.client.post(
                f'/children/{child.id}/exemption/{ticket.id}/use',
                data={'used_on': self.today.isoformat()},
                follow_redirects=True,
            )
            empty = self.client.post(
                f'/children/{child.id}/exemption/{ticket.id}/use',
                data={'subject_key': '', 'used_on': self.today.isoformat()},
                follow_redirects=True,
            )
            english = self.client.post(
                f'/children/{child.id}/exemption/{ticket.id}/use',
                data={'subject_key': 'english', 'used_on': self.today.isoformat()},
                follow_redirects=True,
            )
        html = missing.get_data(as_text=True) + empty.get_data(as_text=True) + english.get_data(as_text=True)
        self.assertIn('면제할 과목을 선택해주세요.', html)
        db.session.refresh(ticket)
        self.assertEqual(ticket.status, 'active')
        self.assertEqual(ExemptionUsage.query.count(), 0)
        self.assertFalse(ExemptionUsage.query.filter_by(subject_key='korean').count())

        spoof_child, spoof_ticket = issue_one('이름조작')
        with mock.patch('features.exemption.service.kst_today', return_value=self.today):
            spoof = self.client.post(
                f'/children/{spoof_child.id}/exemption/{spoof_ticket.id}/use',
                data={
                    'subject_key': 'reading',
                    'subject_name': '해킹과목',
                    'used_on': self.today.isoformat(),
                },
                follow_redirects=False,
            )
        self.assertEqual(spoof.status_code, 302)
        usage = ExemptionUsage.query.filter_by(exemption_ticket_id=spoof_ticket.id).one()
        self.assertEqual(usage.subject_key, 'reading')
        self.assertEqual(usage.subject_name, '독서')
        self.assertNotEqual(usage.subject_name, '해킹과목')

        for key, name in (('korean', '국어'), ('math', '수학'), ('ssen', '쎈'), ('reading', '독서')):
            child, ticket = issue_one(f'허용{name}')
            with mock.patch('features.exemption.service.kst_today', return_value=self.today):
                resp = self.client.post(
                    f'/children/{child.id}/exemption/{ticket.id}/use',
                    data={'subject_key': key, 'used_on': self.today.isoformat()},
                    follow_redirects=False,
                )
            self.assertEqual(resp.status_code, 302)
            db.session.refresh(ticket)
            self.assertEqual(ticket.status, 'used')
            usage = ExemptionUsage.query.filter_by(exemption_ticket_id=ticket.id).one()
            self.assertEqual(usage.subject_key, key)
            self.assertEqual(usage.subject_name, name)

    def test_fixed_exemption_subjects_without_learning_subject_rows(self):
        LearningSubject.query.delete()
        db.session.commit()
        self.assertEqual(LearningSubject.query.count(), 0)
        for key, name in (('korean', '국어'), ('math', '수학'), ('ssen', '쎈'), ('reading', '독서')):
            child = self._child(f'면제{name}', 5)
            self._complete_exemption(child, self._rec_book(), self.today)
            ticket = issue_exemption_ticket(child.id, self.teacher, today=self.today)
            usage = use_exemption_ticket(ticket.id, key, self.teacher, today=self.today)
            self.assertEqual(usage.subject_key, key)
            self.assertEqual(usage.subject_name, name)
        self.assertEqual(LearningSubject.query.count(), 0)
        snap = child_exemption_snapshot(ExemptionTicket.query.first().child_id, self.today)
        keys = [row['key'] for row in snap['eligible_subjects']]
        self.assertEqual(keys, ['korean', 'math', 'ssen', 'reading'])

    def test_expire_then_next_ticket_without_returning_source(self):
        child = self._child('만료반환금지', 5)
        a = self._complete_exemption(child, self._rec_book('만A'), self.today)
        ticket1 = issue_exemption_ticket(child.id, self.teacher, today=self.today)
        self._complete_exemption(child, self._rec_book('만B'), self.today + timedelta(days=1))
        self._complete_exemption(child, self._rec_book('만C'), self.today + timedelta(days=2))
        ticket2 = issue_exemption_ticket(child.id, self.teacher, today=date(2026, 9, 3))
        self.assertEqual(ExemptionTicket.query.get(ticket1.id).status, 'expired')
        self.assertNotIn(
            a.id,
            {
                row.child_reading_id
                for row in ExemptionTicketSource.query.filter_by(exemption_ticket_id=ticket2.id)
            },
        )
        with self.assertRaises(ExemptionError) as err:
            set_reward_mode(ChildReading.query.get(a.id), 'points', self.teacher)
        self.assertEqual(err.exception.code, 'source_expired')
        self.assertNotIn('근거', err.exception.message)
        self.assertIn('면제권 발급에 반영', err.exception.message)

    def test_revoke_returns_source_and_excludes_from_count(self):
        child = self._child('취소반환', 5)
        reading = self._complete_exemption(child, self._rec_book(), self.today)
        ticket = issue_exemption_ticket(child.id, self.teacher, today=self.today)
        revoke_exemption_ticket(ticket.id, self.teacher, today=self.today)
        saved = ExemptionTicket.query.get(ticket.id)
        self.assertEqual(saved.status, 'revoked')
        self.assertIsNotNone(saved.revoked_at)
        snap = child_exemption_snapshot(child.id, self.today)
        self.assertEqual(snap['needed'], 1)
        self.assertTrue(snap['can_issue'])
        self.assertEqual(snap['held_count'], 0)
        again = issue_exemption_ticket(child.id, self.teacher, today=self.today)
        self.assertEqual(
            ExemptionTicketSource.query.filter_by(exemption_ticket_id=again.id).one().child_reading_id,
            reading.id,
        )

    def test_used_ticket_cannot_revoke(self):
        child = self._child('사용취소거부', 5)
        self._complete_exemption(child, self._rec_book(), self.today)
        ticket = issue_exemption_ticket(child.id, self.teacher, today=self.today)
        use_exemption_ticket(ticket.id, 'math', self.teacher, today=self.today)
        with self.assertRaises(ExemptionError) as err:
            revoke_exemption_ticket(ticket.id, self.teacher, today=self.today)
        self.assertEqual(err.exception.code, 'already_used')

    def test_cannot_switch_to_points_after_used_source(self):
        child = self._child('사용전환거부', 5)
        reading = self._complete_exemption(child, self._rec_book(), self.today)
        ticket = issue_exemption_ticket(child.id, self.teacher, today=self.today)
        use_exemption_ticket(ticket.id, 'math', self.teacher, today=self.today)
        with self.assertRaises(ExemptionError) as err:
            set_reward_mode(reading, 'points', self.teacher)
        self.assertEqual(err.exception.code, 'source_used')
        self.assertNotIn('근거', err.exception.message)
        self.assertIn('면제권 발급에 반영되어 포인트 보상으로 변경할 수 없습니다.', err.exception.message)

    def test_issue_and_use_do_not_touch_points_or_progress(self):
        child = self._child('독립', 5)
        save_progress_entry(
            child_id=child.id,
            learning_subject_id=self.math.id,
            textbook_title='보존교재',
            page=12,
            recorded_on=self.today,
            created_by_user_id=self.teacher.id,
        )
        self._complete_exemption(child, self._rec_book(), self.today)
        issue_exemption_ticket(child.id, self.teacher, today=self.today)
        ticket = ExemptionTicket.query.one()
        use_exemption_ticket(ticket.id, 'math', self.teacher, today=self.today)
        self.assertEqual(DailyPoints.query.count(), 0)
        self.assertEqual(PointsHistory.query.count(), 0)
        self.assertEqual(Child.query.get(child.id).cumulative_points or 0, 0)
        self.assertEqual(LearningProgressEntry.query.one().page, 12)

    def test_general_reading_points_coexist_with_exemption_progress(self):
        child = self._child('공존', 5)
        self._complete_exemption(child, self._rec_book(), self.today)
        daily = DailyPoints(
            child_id=child.id,
            date=self.today,
            korean_points=0,
            math_points=0,
            ssen_points=0,
            reading_points=100,
            piano_points=0,
            english_points=0,
            advanced_math_points=0,
            writing_points=0,
            manual_points=0,
            manual_history='[]',
            total_points=100,
            created_by=self.teacher.id,
        )
        db.session.add(daily)
        db.session.commit()
        snap = child_exemption_snapshot(child.id, self.today)
        self.assertTrue(snap['condition_met'])
        self.assertEqual(DailyPoints.query.one().reading_points, 100)
        self.assertEqual(ExemptionTicket.query.count(), 0)

    def test_use_form_lists_fixed_four_subjects(self):
        child = self._child('선택목록', 5)
        self._complete_exemption(child, self._rec_book(), self.today)
        issue_exemption_ticket(child.id, self.teacher, today=self.today)
        self._login(self.teacher_id)
        html = self.client.get(f'/children/{child.id}').get_data(as_text=True)
        self.assertIn('name="subject_key"', html)
        self.assertIn('과목을 선택하세요', html)
        self.assertIn('value="" selected', html)
        self.assertNotIn('value="korean" selected', html)
        self.assertIn('value="korean"', html)
        self.assertIn('value="math"', html)
        self.assertIn('value="ssen"', html)
        self.assertIn('value="reading"', html)
        self.assertIn('>국어<', html)
        self.assertIn('>수학<', html)
        self.assertIn('>쎈<', html)
        self.assertIn('>독서<', html)
        self.assertIn('data-exemption-use-submit', html)
        self.assertRegex(html, r'data-exemption-use-submit[^>]*\bdisabled\b')

    def test_viewer_cannot_manage_learning_subjects(self):
        self._login(self.viewer_id)
        blocked = self.client.get('/settings/learning-subjects', follow_redirects=False)
        self.assertEqual(blocked.status_code, 302)
        post_resp = self.client.post(
            '/settings/learning-subjects',
            data={'key': 'hack', 'name': '해킹'},
            follow_redirects=False,
        )
        self.assertEqual(post_resp.status_code, 302)
        self.assertIsNone(LearningSubject.query.filter_by(key='hack').first())

    def test_learning_subject_inactive_does_not_change_exemption_choices(self):
        set_subject_active(self.ssen, False)
        snap = child_exemption_snapshot(self._child('선택목록', 5).id, self.today)
        keys = [row['key'] for row in snap['eligible_subjects']]
        self.assertEqual(keys, ['korean', 'math', 'ssen', 'reading'])

    def test_reset_preserves_ledger_and_first_ticket_count(self):
        child = self._child('학기보존', 5)
        self._complete_exemption(child, self._rec_book(), self.today)
        issue_exemption_ticket(child.id, self.teacher, today=self.today)
        self._login(self.developer_id)
        resp = self.client.post('/settings/data', data={'action': 'reset_data'}, follow_redirects=False)
        self.assertIn(resp.status_code, (200, 302))
        self.assertEqual(ExemptionTicket.query.count(), 1)
        self.assertEqual(ExemptionTicketSource.query.count(), 1)
        self.assertEqual(ChildReading.query.one().reward_mode, 'exemption')
        snap = child_exemption_snapshot(child.id, self.today)
        self.assertEqual(snap['needed'], 2)
        self.assertEqual(DailyPoints.query.count(), 0)

    def test_json_backup_restore_and_excel_sheets(self):
        child = self._child('백업면제', 5)
        reading = self._complete_exemption(child, self._rec_book(), self.today)
        ticket = issue_exemption_ticket(child.id, self.teacher, today=self.today)
        use_exemption_ticket(ticket.id, 'math', self.teacher, today=self.today)
        backup_data, error = get_backup_data()
        self.assertIsNone(error)
        counts = backup_data['backup_metadata']['records_count']
        self.assertEqual(counts['exemption_tickets'], 1)
        self.assertEqual(counts['exemption_ticket_sources'], 1)
        self.assertEqual(counts['exemption_usages'], 1)
        self.assertEqual(backup_data['child_readings'][0]['reward_mode'], 'exemption')
        self.assertTrue(any(row['key'] == 'math' for row in backup_data['learning_subjects']))
        self.assertNotIn('is_exemption_eligible', backup_data['learning_subjects'][0])
        self.assertEqual(backup_data['exemption_usages'][0]['subject_key'], 'math')
        self.assertEqual(backup_data['exemption_usages'][0]['subject_name'], '수학')

        ExemptionUsage.query.delete()
        ExemptionTicketSource.query.delete()
        ExemptionTicket.query.delete()
        ReadingDay.query.delete()
        ChildReading.query.delete()
        Book.query.delete()
        db.session.commit()
        restore_books_from_backup_data(backup_data)
        restore_readings_from_backup_data(backup_data)
        restore_exemptions_from_backup_data(backup_data)
        restored = ExemptionTicket.query.one()
        self.assertEqual(restored.status, 'used')
        self.assertEqual(ExemptionTicketSource.query.one().child_reading_id, reading.id)
        self.assertEqual(ExemptionUsage.query.one().subject_key, 'math')
        self.assertEqual(ExemptionUsage.query.one().subject_name, '수학')
        self.assertEqual(ChildReading.query.one().reward_mode, 'exemption')

        from app import BACKUP_EXCEL_AVAILABLE, create_excel_backup
        if BACKUP_EXCEL_AVAILABLE:
            with tempfile.TemporaryDirectory() as tmp:
                (Path(tmp) / 'realtime').mkdir(parents=True, exist_ok=True)
                path, excel_error = create_excel_backup(backup_data, tmp, backup_type='manual')
                self.assertIsNone(excel_error)
                from openpyxl import load_workbook
                wb = load_workbook(path)
                self.assertIn('면제권', wb.sheetnames)
                self.assertIn('면제권출처', wb.sheetnames)
                self.assertIn('면제권사용이력', wb.sheetnames)

    def test_teacher_http_issue_and_use(self):
        child = self._child('화면발급', 5)
        self._complete_exemption(child, self._rec_book(), self.today)
        self._login(self.teacher_id)
        with mock.patch('features.exemption.service.kst_today', return_value=self.today):
            issue = self.client.post(f'/children/{child.id}/exemption/issue', follow_redirects=False)
            self.assertEqual(issue.status_code, 302)
            ticket = ExemptionTicket.query.one()
            editor = self.client.get(f'/children/{child.id}')
            html = editor.get_data(as_text=True)
            self.assertIn('면제권 1장 보유', html)
            self.assertIn('면제권 사용 · 종이 회수', html)
            use = self.client.post(
                f'/children/{child.id}/exemption/{ticket.id}/use',
                data={'subject_key': 'reading', 'used_on': self.today.isoformat()},
                follow_redirects=False,
            )
            self.assertEqual(use.status_code, 302)
        self.assertEqual(ExemptionTicket.query.one().status, 'used')
        self.assertEqual(ExemptionUsage.query.one().subject_key, 'reading')
        self.assertEqual(ExemptionUsage.query.one().subject_name, '독서')

    def test_same_day_second_recommended_is_blocked(self):
        child = self._child('같은날두번째', 5)
        self._complete_exemption(child, self._rec_book('추천A'), self.today)
        with self.assertRaises(ReadingError) as err:
            start_book(
                child.id,
                self.teacher.id,
                'teacher',
                activity_date=self.today,
                book_id=self._rec_book('추천B').id,
            )
        self.assertEqual(err.exception.code, 'already_logged_today')

    def test_grade_2_4_unchanged_no_mode(self):
        child = self._child('4학년유지', 4)
        book, _ = create_book_record('4학년추천', '작가')
        update_recommended_flags(book, is_recommended=True, grade_band='4-6')
        reading, _ = self._start(child, book, self.today)
        with self.assertRaises(ExemptionError) as err:
            set_reward_mode(reading, 'exemption', self.teacher)
        self.assertEqual(err.exception.code, 'not_eligible')
        approve_recommended_reward(reading, self.teacher, EVENT_RECOMMENDED_START)
        self.assertEqual(DailyPoints.query.one().manual_points, 100)


if __name__ == '__main__':
    unittest.main()
